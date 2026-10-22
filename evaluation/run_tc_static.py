#!/usr/bin/env python3
"""
Static checks that need no server and no LLM quota (QA process groups 1, 2, 5).

  Group 1 Ingestion/OCR : pages per PDF vs pages covered by chunks, broken
                          Vietnamese fragments, odd characters, junk chunks,
                          lines repeated across many chunks (watermarks).
  Group 2 Chunking      : empty metadata, page range, token size, chunks that
                          start mid-sentence; samples for manual reading.
  Group 5 Citation      : from <run-dir>/raw.jsonl - answered cases with no
                          citation, combined tags [SOURCE_2, SOURCE_4], tags in
                          the answer that the API dropped, cited doc vs expected.

    python evaluation/run_tc_static.py --input ~/Downloads/TC_LuatLD.xlsx \
        --run-dir evaluation/results/2026-10-20_flash-lite
Writes <run-dir>/static_checks.json and <run-dir>/static_checks.xlsx (samples
for the tester to fill in).
"""
import argparse
import json
import random
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))

from tc_common import load_cases, same_doc  # noqa: E402

HARD_MAX_TOKENS = 800          # src/data/block_packer.py
MIN_BODY_CHARS = 40            # shorter bodies are junk candidates
MIN_PAGE_CHARS = 60            # PDF pages with less text are blank / header-only
WATERMARK_SHARE = 0.05         # a line in >= 5% of a doc's chunks
FRAGMENT_LIMIT = 50.0          # same threshold as src/data/pdf_loader.py
VOWELS = set("aăâeêioôơuưyáàảãạắằẳẵặấầẩẫậéèẻẽẹếềểễệíìỉĩịóòỏõọốồổỗộớờởỡợúùủũụứừửữựýỳỷỹỵ")
REFUSAL = "không tìm thấy thông tin để trả lời"
COMBINED_TAG = re.compile(r"\[SOURCE_\d+(?:\s*[,;]\s*(?:SOURCE_)?\d+)+\]")
SINGLE_TAG = re.compile(r"\[SOURCE_(\d+)\]")
LIST_MARKER = re.compile(r"^[a-zđ]\d*\s*[\.\)\-]")
ANY_TAG_ID = re.compile(r"SOURCE_(\d+)")


def fragment_rate(text):
    # Mirrors src/data/pdf_loader.fragment_rate without importing docling.
    words = re.findall(r"[^\W\d_]+", unicodedata.normalize("NFC", text or ""))
    if not words:
        return 0.0
    broken = sum(1 for w in words if w.islower() and len(w) <= 3 and not (set(w) & VOWELS))
    return 1000.0 * broken / len(words)


def odd_chars(text):
    return sum(1 for ch in text if ch == "�" or unicodedata.category(ch) in ("Co", "Cc") and ch not in "\n\t")


def body(chunk):
    return chunk["chunk_content"].split("\n---\n", 1)[-1]


def check_ingestion(chunks, raw_dir):
    import pypdfium2

    by_doc = defaultdict(list)
    for c in chunks:
        by_doc[c["metadata"]["source"]].append(c)
    docs = {}
    for pdf in sorted(raw_dir.glob("*.pdf")):
        name = unicodedata.normalize("NFC", pdf.name)
        cs = next((v for k, v in by_doc.items() if same_doc(k, name)), [])
        n_pages = len(pypdfium2.PdfDocument(str(pdf)))
        covered = set()
        for c in cs:
            m = c["metadata"]
            covered.update(range(int(m.get("page_start") or 0), int(m.get("page_end") or 0) + 1))
        # Blank pages (cover, separators) and pages with only a running header are not "missing".
        pdf = pypdfium2.PdfDocument(str(pdf))
        missing = sorted(p for p in set(range(1, n_pages + 1)) - covered
                         if len(re.sub(r"\W|\d", "", pdf[p - 1].get_textpage().get_text_range())) > MIN_PAGE_CHARS)
        text = "\n".join(body(c) for c in cs)
        lines = Counter()
        for c in cs:
            lines.update({ln.strip() for ln in body(c).splitlines() if len(ln.strip()) >= 8})
        watermark = [(ln, n) for ln, n in lines.most_common(10) if n >= max(5, WATERMARK_SHARE * len(cs))]
        junk = [c["chunk_id"] for c in cs if len(re.sub(r"\W", "", body(c))) < MIN_BODY_CHARS]
        rate = fragment_rate(text)
        docs[name] = {
            "pdf_pages": n_pages, "chunks": len(cs), "pages_covered": len(covered & set(range(1, n_pages + 1))),
            "pages_missing": len(missing), "pages_missing_list": missing[:40],
            "fragment_rate": round(rate, 1), "fragment_ok": rate <= FRAGMENT_LIMIT,
            "odd_chars": odd_chars(text), "junk_chunks": len(junk), "junk_examples": junk[:5],
            "repeated_lines": [{"line": ln[:120], "chunks": n} for ln, n in watermark],
        }
    return docs


def check_chunking(chunks):
    issues = defaultdict(list)
    for c in chunks:
        m, cid = c["metadata"], c["chunk_id"]
        for key in ("source", "section", "page_start", "page_end"):
            if m.get(key) in (None, "", 0):
                issues[f"empty_{key}"].append(cid)
        if int(m.get("page_start") or 0) > int(m.get("page_end") or 0):
            issues["page_start_gt_end"].append(cid)
        if int(m.get("token_count") or 0) > HARD_MAX_TOKENS:
            issues["over_hard_max"].append(cid)
        # Lowercase start that is not a list marker ("b - ...", "d1) ...", "c. ...").
        # Expected for the overlap tail carried from the previous chunk, so it is informational.
        head = body(c).lstrip("#*-•· \n")
        if head[:1].islower() and not LIST_MARKER.match(head):
            issues["starts_mid_sentence"].append(cid)
    tokens = sorted(int(c["metadata"].get("token_count") or 0) for c in chunks)
    return {
        "chunks": len(chunks), "unique_ids": len({c["chunk_id"] for c in chunks}),
        "tokens_p50": tokens[len(tokens) // 2], "tokens_max": tokens[-1],
        "issues": {k: {"count": len(v), "examples": v[:5]} for k, v in sorted(issues.items())},
    }


def check_citations(cases, run_dir):
    raw = {}
    for line in (run_dir / "raw.jsonl").read_text(encoding="utf-8").splitlines():
        if line.strip():
            rec = json.loads(line)
            raw[rec["id"]] = rec
    rows = []
    for case in cases:
        resp = (raw.get(case["id"]) or {}).get("response")
        if not resp:
            continue
        answer = resp.get("answer") or ""
        if REFUSAL in answer.lower():
            continue
        cited = resp.get("citations") or []
        sent = {s["source_id"] for s in (resp.get("debug") or {}).get("context_sources") or []}
        returned = {c.get("source_id") for c in cited}
        in_text = {f"SOURCE_{n}" for n in ANY_TAG_ID.findall(answer)}
        combined = COMBINED_TAG.findall(answer)
        rows.append({
            "id": case["id"], "group": case["group"], "n_citations": len(cited),
            "combined_tags": combined, "single_tags": len(SINGLE_TAG.findall(answer)),
            # Tags the model wrote and that were in the context, but the API returned no citation for.
            "dropped": sorted((in_text & sent) - returned),
            "invented": sorted(in_text - sent),
            "cites_expected_doc": any(same_doc(c.get("file"), case["expected_doc"]) for c in cited),
        })
    n = len(rows)
    return {
        "answered": n,
        "with_citation": sum(r["n_citations"] > 0 for r in rows),
        "no_citation": [r["id"] for r in rows if not r["n_citations"]],
        "combined_tag_cases": [r["id"] for r in rows if r["combined_tags"]],
        "dropped_tag_cases": [r["id"] for r in rows if r["dropped"]],
        "invented_tag_cases": [r["id"] for r in rows if r["invented"]],
        "cites_expected_doc": sum(r["cites_expected_doc"] for r in rows),
        "rows": rows,
    }, raw


def write_workbook(path, chunks, cases, raw, seed):
    import openpyxl
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    rnd = random.Random(seed)
    wb = openpyxl.Workbook()

    def sheet(title, header, rows, widths, verdict_col):
        ws = wb.create_sheet(title)
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        dv = DataValidation(type="list", formula1='"OK,Lỗi,Không chắc"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{verdict_col}2:{verdict_col}{len(rows) + 1}")
        ws.freeze_panes = "A2"

    # Group 1: 30 random chunks per document, to compare with the PDF page.
    by_doc = defaultdict(list)
    for c in chunks:
        by_doc[c["metadata"]["source"]].append(c)
    rows = []
    for doc, cs in sorted(by_doc.items()):
        for c in rnd.sample(cs, min(30, len(cs))):
            m = c["metadata"]
            rows.append([doc, m.get("page_start"), m.get("page_end"), c["chunk_id"], body(c)[:1500], "", ""])
    sheet("G1_OCR_Sample", ["Document", "Page start", "Page end", "Chunk ID", "Text (compare with PDF)", "Verdict", "Note"],
          rows, [30, 8, 8, 30, 100, 12, 30], "F")

    # Group 2: 20 chunks that contain a legal article heading.
    art = [c for c in chunks if re.search(r"\bĐiều\s+\d+", body(c))]
    rows = [[c["metadata"]["source"], c["metadata"].get("page_start"), c["metadata"].get("section"), c["chunk_id"],
             body(c)[:1500], "", ""] for c in rnd.sample(art, min(20, len(art)))]
    sheet("G2_Article_Sample", ["Document", "Page", "Section", "Chunk ID", "Text (article cut mid-way? metadata right?)", "Verdict", "Note"],
          rows, [30, 8, 40, 30, 100, 12, 30], "F")

    # Group 5: 30 answered cases, each citation to open on the UI / in the PDF.
    answered = [c for c in cases if (raw.get(c["id"]) or {}).get("response")
                and REFUSAL not in ((raw[c["id"]]["response"].get("answer") or "").lower())]
    rows = []
    for case in sorted(rnd.sample(answered, min(30, len(answered))), key=lambda c: c["id"]):
        resp = raw[case["id"]]["response"]
        cites = "\n".join(f"{c.get('source_id')}: {c.get('file')} p.{c.get('page')}" for c in resp.get("citations") or [])
        rows.append([case["id"], case["query"], (resp.get("answer") or "")[:2000], cites or "(none)", "", ""])
    sheet("G5_Citation_Sample", ["Test Case ID", "Query", "Answer", "Citations (open each one)", "Verdict", "Note"],
          rows, [14, 40, 80, 45, 12, 30], "E")

    del wb["Sheet"]
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", type=Path, required=True)
    ap.add_argument("--run-dir", type=Path, required=True)
    ap.add_argument("--chunked", type=Path, default=ROOT / "data" / "processed" / "chunked.jsonl")
    ap.add_argument("--raw-dir", type=Path, default=ROOT / "data" / "raw")
    ap.add_argument("--seed", type=int, default=20261020)
    args = ap.parse_args()

    chunks = [json.loads(line) for line in args.chunked.read_text(encoding="utf-8").splitlines() if line.strip()]
    _, cases = load_cases(args.input.expanduser())
    ingestion = check_ingestion(chunks, args.raw_dir)
    chunking = check_chunking(chunks)
    citations, raw = check_citations(cases, args.run_dir)

    out = args.run_dir / "static_checks.json"
    out.write_text(json.dumps({"ingestion": ingestion, "chunking": chunking, "citations": citations},
                              ensure_ascii=False, indent=1), encoding="utf-8")
    write_workbook(args.run_dir / "static_checks.xlsx", chunks, cases, raw, args.seed)

    print("== Group 1: ingestion")
    for doc, d in ingestion.items():
        print(f"  {doc[:45]:45} pages {d['pages_covered']}/{d['pdf_pages']} (missing text pages {d['pages_missing']}) chunks {d['chunks']:4} "
              f"frag {d['fragment_rate']:5} odd {d['odd_chars']:3} junk {d['junk_chunks']:3} repeated {len(d['repeated_lines'])}")
        for r in d["repeated_lines"][:3]:
            print(f"      x{r['chunks']}: {r['line'][:90]}")
    print("== Group 2: chunking", {k: v for k, v in chunking.items() if k != "issues"})
    for k, v in chunking["issues"].items():
        print(f"  {k}: {v['count']}")
    c = citations
    print(f"== Group 5: citations - answered {c['answered']}, with citation {c['with_citation']}, "
          f"cites expected doc {c['cites_expected_doc']}")
    for k in ("no_citation", "combined_tag_cases", "dropped_tag_cases", "invented_tag_cases"):
        print(f"  {k}: {len(c[k])} {c[k]}")
    print("->", out, "and static_checks.xlsx")


if __name__ == "__main__":
    main()
