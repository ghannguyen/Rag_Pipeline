#!/usr/bin/env python3
"""Đánh giá RAG bằng Gemini, không cần OpenAI API key.

  pip install -r requirements.txt
  export GOOGLE_API_KEY='...'
  python evaluate_rag.py --input 'RAG_TestCase.xlsx' --limit 3
  python evaluate_rag.py --input 'RAG_TestCase.xlsx' --output full_results

Đọc các sheet của workbook, tự tìm header trong 20 dòng đầu.
Ghi CSV/dashboard sau mỗi case. --limit chọn các case hợp lệ theo thứ tự workbook.
Không dùng Evidence làm context; context phải lấy từ API retrieval.
Schema tham chiếu: Desktop/RagPipeline/src/app.py, kiểm tra ngày 2026-09-23.
Live /docs chưa truy cập được lúc viết; luôn kiểm tra OpenAPI khi chạy.
Ragas: https://docs.ragas.io/en/v0.4.2/extra/components/choose_evaluator_llm/
"""
from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import hashlib
import html
import json
import math
import os
from pathlib import Path
import re
import sys
import time

os.environ.setdefault('RAGAS_DO_NOT_TRACK', 'true')
METRICS = ('faithfulness', 'answer_relevancy', 'context_precision', 'context_recall')
SHEETS = ('RAG_TC', 'RAG_TC2')
CONTEXT_NOTE = ('Contexts lấy từ search trước rerank, không đảm bảo trùng context '
                'chat thực sự sử dụng. Faithfulness đo với tập context này. '
                'Bốn metric không phải bộ chấm riêng cho hành vi từ chối trả lời.')


def text(value):
    return '' if value is None else str(value).strip()


def load_cases(path: Path, limit: int | None = None):
    if path.suffix.lower() == '.json':
        records = json.loads(path.read_text(encoding='utf-8'))
        if not isinstance(records, list) or not records:
            raise ValueError('JSON testcase phải là list không rỗng')
        cases = []
        for i, record in enumerate(records[:limit] if limit else records, 1):
            question = text(record.get('query'))
            reference = text(record.get('expected_answer'))
            cases.append(dict(sheet=path.name, excel_row=i,
                test_case_id=text(record.get('id', i)), category='',
                question=question, ground_truth=reference, evidence='',
                is_answerable='Unknown',
                input_error='' if question and reference else 'Thiếu query hoặc expected_answer'))
        return cases
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    cases = []
    try:
        for name in wb.sheetnames:
            headers = None
            seen = set()
            for line, values in enumerate(wb[name].iter_rows(values_only=True), 1):
                values = list(values)
                if headers is None:
                    labels = [text(v) for v in values]
                    required = {'Test Case ID', 'Query', 'Expected Answer'}
                    if required <= set(labels):
                        headers = {v: i for i, v in enumerate(labels) if v}
                    elif line >= 20:
                        raise ValueError(f'{name}: không tìm thấy header trong 20 dòng đầu')
                    continue
                def get(key):
                    i = headers.get(key)
                    return text(values[i]) if i is not None and i < len(values) else ''
                if not any(get(k) for k in ('Test Case ID', 'Query', 'Expected Answer')):
                    continue
                case_id = get('Test Case ID')
                error = ''
                if not case_id or not get('Query') or not get('Expected Answer'):
                    error = 'Thiếu Test Case ID, Query hoặc Expected Answer'
                if case_id in seen:
                    error = 'Test Case ID trùng trong cùng sheet'
                seen.add(case_id)
                flag = get('Is Answerable').casefold()
                # Sheet thiếu nhãn giữ Unknown, không tự suy đoán từ Category.
                answerable = {'yes': 'Yes', 'no': 'No', '': 'Unknown'}.get(flag, 'Unknown')
                if flag not in ('yes', 'no', ''):
                    error = 'Is Answerable phải là Yes/No hoặc trống'
                cases.append(dict(sheet=name, excel_row=line, test_case_id=case_id,
                    category=get('Category'), question=get('Query'),
                    ground_truth=get('Expected Answer'), evidence=get('Evidence'),
                    expected_page=get('Source Page'), expected_section=get('Source Section'),
                    difficulty=get('Retrieval Difficulty'), retrieval_type=get('Expected Retrieval Type'),
                    excel_fields={text(label) or f'unnamed_column_{i+1}': values[i] if i < len(values) else None
                                  for i, label in enumerate(labels)},
                    is_answerable=answerable, input_error=error))
            if headers is None:
                raise ValueError(f'{name}: sheet rỗng hoặc thiếu header')
    finally:
        wb.close()
    if not cases:
        raise ValueError('Không có testcase')
    valid_cases = [case for case in cases if not case['input_error']]
    return valid_cases[:limit] if limit else cases


def resolve_schema(schema, root):
    # Chỉ resolve ref nội bộ; không tải schema ngoài.
    if isinstance(schema, list):
        return [resolve_schema(s, root) for s in schema]
    if not isinstance(schema, dict):
        return schema
    if '$ref' in schema:
        ref = schema['$ref']
        if not ref.startswith('#/'):
            raise ValueError('OpenAPI chứa reference ngoài, cần kiểm tra thủ công')
        target = root
        for key in ref[2:].split('/'):
            target = target[key.replace('~1', '/').replace('~0', '~')]
        return resolve_schema(target, root)
    return {k: resolve_schema(v, root) for k, v in schema.items()}


def check_openapi(session, args):
    import jsonschema
    url = args.base_url.rstrip('/') + '/openapi.json'
    response = session.get(url, timeout=(10, 20))
    response.raise_for_status()
    doc = response.json()
    bodies = {
        '/api/rag/chat': dict(query='schema check', top_k=args.top_k,
            top_n=args.top_n, threshold=args.threshold, provider='gemini'),
        '/api/retrieval/search': dict(query='schema check', top_k=args.top_k),
    }
    for path, body in bodies.items():
        try:
            operation = doc['paths'][path]['post']
            request = resolve_schema(operation['requestBody'], doc)
            schema = resolve_schema(request['content']['application/json']['schema'], doc)
        except KeyError as exc:
            raise ValueError(f'OpenAPI thiếu POST/JSON schema cho {path}') from exc
        properties = schema.get('properties', {})
        if set(body) - set(properties):
            raise ValueError(f'{path}: field không khớp OpenAPI: {set(body) - set(properties)}')
        jsonschema.validate(body, schema)
        for parameter in operation.get('parameters', []):
            if parameter.get('required'):
                raise ValueError(f'{path}: có parameter bắt buộc chưa cấu hình: {parameter.get("name")}')
    return hashlib.sha256(json.dumps(doc, sort_keys=True).encode()).hexdigest()


def call_json(session, url, body, timeout):
    # Không tự retry POST dài để tránh chạy lại pipeline và tính phí hai lần.
    response = session.post(url, json=body, timeout=(10, timeout))
    response.raise_for_status()
    return response.json()


def parse_answer(payload):
    # Source backend trả object.answer; vẫn nhận JSON string như mô tả API.
    answer = payload.get('answer') if isinstance(payload, dict) else payload
    if not isinstance(answer, str) or not answer.strip():
        raise ValueError('Chat phải trả JSON string hoặc object.answer không rỗng')
    return answer


def parse_contexts(payload):
    if not isinstance(payload, dict) or not isinstance(payload.get('results'), list):
        raise ValueError('Retrieval phải trả object.results là list')
    contexts = []
    for item in payload['results']:
        if not isinstance(item, dict) or not isinstance(item.get('content'), str):
            raise ValueError('Mỗi retrieval result phải có content dạng string')
        if not item['content'].strip():
            raise ValueError('Retrieval trả chunk content rỗng')
        contexts.append(item['content'])
    return contexts


def safe_error(exc):
    message = str(exc)
    for key in ('GOOGLE_API_KEY', 'GEMINI_API_KEY'):
        if os.getenv(key):
            message = message.replace(os.environ[key], '[REDACTED]')
    message = re.sub(r'AIza[\w-]+', '[REDACTED]', message)
    return f'{type(exc).__name__}: {message[:1500]}'


def make_evaluator(args):
    from langchain_google_genai import ChatGoogleGenerativeAI, GoogleGenerativeAIEmbeddings
    from langchain_core.rate_limiters import InMemoryRateLimiter
    from ragas.llms import LangchainLLMWrapper
    from ragas.embeddings import LangchainEmbeddingsWrapper
    from ragas.run_config import RunConfig
    key = os.getenv('GOOGLE_API_KEY') or os.getenv('GEMINI_API_KEY')
    if not key:
        raise ValueError('Cần GOOGLE_API_KEY hoặc GEMINI_API_KEY trong môi trường/.env')
    config = RunConfig(timeout=args.judge_timeout, max_retries=2, max_wait=20, max_workers=1)
    rate_limiter = InMemoryRateLimiter(
        requests_per_second=1 / 13,
        check_every_n_seconds=0.5,
        max_bucket_size=1,
    )

    llm = LangchainLLMWrapper(ChatGoogleGenerativeAI(
        model=args.judge_model,
        google_api_key=key,
        temperature=0,
        max_output_tokens=8192,
        timeout=args.judge_timeout,
        max_retries=1,
        rate_limiter=rate_limiter,
    ), run_config=config)
    # Cùng task SEMANTIC_SIMILARITY cho câu hỏi gốc và câu hỏi tái tạo.
    embeddings = LangchainEmbeddingsWrapper(GoogleGenerativeAIEmbeddings(
        model=args.embedding_model, google_api_key=key,
        task_type='SEMANTIC_SIMILARITY'), run_config=config)
    return llm, embeddings, config


def score_case(row, evaluator):
    from ragas import evaluate, EvaluationDataset
    from ragas.metrics import Faithfulness, ResponseRelevancy, LLMContextPrecisionWithReference, LLMContextRecall
    from langchain_core.callbacks import BaseCallbackHandler

    class Errors(BaseCallbackHandler):
        def __init__(self):
            self.messages = []
        def on_chain_error(self, error, **kwargs):
            self.messages.append(safe_error(error))
        def on_llm_error(self, error, **kwargs):
            self.messages.append(safe_error(error))

    llm, embeddings, config = evaluator
    metrics = [
        Faithfulness(),
        ResponseRelevancy(strictness=1),
        LLMContextPrecisionWithReference(),
        LLMContextRecall(),
    ]
    data = EvaluationDataset.from_list([dict(user_input=row['question'],
        reference=row['ground_truth'], response=row['answer'], retrieved_contexts=row['contexts'])])
    errors = Errors()
    result = evaluate(data, metrics=metrics, llm=llm, embeddings=embeddings,
        run_config=config, callbacks=[errors], raise_exceptions=False, show_progress=False)
    scores = result.scores[0]
    missing = []
    for name, metric in zip(METRICS, metrics):
        value = scores.get(metric.name)
        if value is None or not math.isfinite(float(value)):
            row[name] = None
            missing.append(name)
        else:
            row[name] = float(value)
    if missing:
        row['error'] = 'Metric không có điểm: ' + ', '.join(missing)
        if errors.messages:
            row['error'] += ' | ' + ' | '.join(dict.fromkeys(errors.messages))
        row['status'] = 'METRIC_ERROR'
    else:
        row['status'] = 'OK'


def atomic_text(path, content):
    temp = path.with_suffix(path.suffix + '.tmp')
    temp.write_text(content, encoding='utf-8')
    temp.replace(path)


def save_outputs(rows, args, metadata):
    fields = ['sheet', 'excel_row', 'test_case_id', 'category', 'is_answerable',
        'question', 'ground_truth', 'evidence', 'answer', 'contexts', *METRICS,
        'status', 'error', 'retrieval_seconds', 'chat_seconds', 'judge_seconds',
        'total_seconds', 'context_source', 'judge_model', 'embedding_model', 'run_started']
    target = args.output / 'results_raw.csv'
    temp = target.with_suffix('.csv.tmp')
    with temp.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction='ignore')
        writer.writeheader()
        for row in rows:
            value = {**row, 'contexts': json.dumps(row['contexts'], ensure_ascii=False),
                     'judge_model': args.judge_model, 'embedding_model': args.embedding_model,
                     'run_started': metadata['started']}
            writer.writerow(value)
    temp.replace(target)
    # HTML độc lập, SVG inline; mở offline, không cần CDN.
    esc = lambda s: html.escape(str(s), quote=True)
    cards, charts = [], []
    for metric in METRICS:
        valid = [r[metric] for r in rows if r.get(metric) is not None]
        average = sum(valid) / len(valid) if valid else None
        mean = f'{average:.3f}' if average is not None else '—'
        cards.append(f'<article><label>{metric}</label><strong>{mean}</strong><small>{len(valid)}/{len(rows)} có điểm</small></article>')
        bars = []
        bins = [0] * 10
        for value in valid:
            bins[min(9, max(0, int(value * 10)))] += 1
        maximum = max(bins) or 1
        for i, count in enumerate(bins):
            height = 130 * count / maximum
            bars.append(f'<rect x="{35+i*36}" y="{155-height}" width="26" height="{height}" rx="3"><title>{i/10:.1f}–{(i+1)/10:.1f}: {count} case</title></rect><text x="{48+i*36}" y="{148-height}" text-anchor="middle">{count}</text>')
        charts.append(f'<article><h3>{metric}</h3><svg role="img" aria-label="Phân bố điểm {metric}" viewBox="0 0 420 190">{"".join(bars)}<text x="35" y="180">0</text><text x="370" y="180">1.0</text></svg></article>')
    groups = []
    for group in ('Yes', 'No', 'Unknown'):
        subset = [r for r in rows if r['is_answerable'] == group]
        values = []
        for metric in METRICS:
            numbers = [r[metric] for r in subset if r.get(metric) is not None]
            values.append(f'{sum(numbers)/len(numbers):.3f} (n={len(numbers)})' if numbers else '—')
        groups.append('<tr>' + ''.join(f'<td>{esc(v)}</td>' for v in [group, len(subset), *values]) + '</tr>')
    detail = []
    for row in rows:
        scores = [f'{row[m]:.3f}' if row.get(m) is not None else '—' for m in METRICS]
        cells = [row['sheet'], row['test_case_id'], row['is_answerable'], row['status'], *scores]
        detail.append('<tr>' + ''.join(f'<td>{esc(v)}</td>' for v in cells) +
            '<td><details><summary>' + esc(row['question']) + '</summary><h4>Answer</h4><pre>' +
            esc(row['answer']) + '</pre><h4>Ground truth</h4><pre>' + esc(row['ground_truth']) +
            '</pre><h4>Contexts</h4><pre>' + esc('\n\n'.join(row['contexts'])) + '</pre><p>' + esc(row['error']) + '</p></details></td></tr>')
    ok = sum(r['status'] == 'OK' for r in rows)
    document = '''<!doctype html><html lang="vi"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>RAG evaluation</title><style>
body{font:15px system-ui,sans-serif;background:#f4f6fa;color:#17243c;margin:0;padding:32px}main{max-width:1440px;margin:auto}h1{margin:0;font-size:30px}p{line-height:1.6;color:#516077}.grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(260px,1fr));gap:16px;margin:24px 0}article{background:white;border:1px solid #dce3ee;border-radius:12px;padding:22px}strong{display:block;font-size:34px;margin:12px 0}small,label{display:block;color:#516077}svg{width:100%;fill:#4576df}svg text{fill:#516077;font-size:11px}table{border-collapse:collapse;width:100%;background:white}td,th{padding:12px;text-align:left;border-bottom:1px solid #e1e7f0;vertical-align:top}th{font-size:12px;color:#516077}.scroll{overflow:auto;border-radius:10px;margin:20px 0}pre{white-space:pre-wrap;overflow-wrap:anywhere;font:13px system-ui;max-width:700px}summary{cursor:pointer;min-width:260px}h3{font-size:16px}input{padding:12px;border:1px solid #ccd5e5;border-radius:8px;width:min(400px,90%)}.note{background:#fff5de;padding:16px;border-radius:10px}
</style><main>'''
    document += f'<h1>RAG evaluation</h1><p>{esc(metadata["started"])} · {esc(args.judge_model)} · {len(rows)}/{metadata["selected"]} testcase đã xử lý · {ok} đủ 4 điểm · {len(rows)-ok} có lỗi</p>'
    document += f'<p class="note">{esc(CONTEXT_NOTE)} Điểm lỗi để trống, không đổi thành 0. Trung bình chỉ tính điểm hợp lệ; OK nghĩa là chấm thành công, không có nghĩa PASS.</p>'
    document += '<div class="grid">' + ''.join(cards) + '</div><div class="grid">' + ''.join(charts) + '</div>'
    document += '<h2>Theo Is Answerable</h2><p>Unknown: sheet không có nhãn hoặc ô trống. Nhóm No cần đọc riêng vì đáp án kỳ vọng là từ chối.</p><div class="scroll"><table><thead><tr>'
    document += ''.join(f'<th>{m}</th>' for m in ('Is Answerable', 'Cases', *METRICS)) + '</tr></thead><tbody>' + ''.join(groups) + '</tbody></table></div>'
    document += '<h2>Chi tiết testcase</h2><input id="filter" placeholder="Tìm case, câu hỏi hoặc lỗi"><div class="scroll"><table id="details"><thead><tr>'
    document += ''.join(f'<th>{m}</th>' for m in ('Sheet', 'ID', 'Answerable', 'Status', *METRICS, 'Nội dung')) + '</tr></thead><tbody>' + ''.join(detail) + '</tbody></table></div>'
    document += '<p>OpenAPI SHA256: ' + esc(metadata.get('schema_hash', 'chưa kết nối')) + '</p></main><script>document.getElementById("filter").addEventListener("input",function(){const q=this.value.toLowerCase();document.querySelectorAll("#details tbody tr").forEach(r=>r.hidden=!r.textContent.toLowerCase().includes(q));});</script></html>'
    atomic_text(args.output / 'dashboard.html', document)


def run(args):
    import requests
    from dotenv import load_dotenv
    load_dotenv(args.env_file or Path.cwd() / '.env', override=False)
    cases = load_cases(args.input, args.limit)
    args.output.mkdir(parents=True, exist_ok=True)
    if not args.overwrite and any((args.output / name).exists() for name in ('results_raw.csv', 'dashboard.html')):
        raise ValueError('Output đã tồn tại; chọn --output khác hoặc dùng --overwrite')
    metadata = dict(started=datetime.now(timezone.utc).isoformat(), selected=len(cases))
    rows = []
    with requests.Session() as session:
        # Không gửi localhost qua proxy môi trường.
        session.trust_env = False
        print('Kiểm tra /openapi.json...', flush=True)
        metadata['schema_hash'] = check_openapi(session, args)
        evaluator = make_evaluator(args)
        save_outputs(rows, args, metadata)
        for i, case in enumerate(cases, 1):
            row = {**case, 'answer': '', 'contexts': [], **dict.fromkeys(METRICS),
                   'status': 'ERROR', 'error': '', 'context_source': 'retrieval_search_pre_rerank',
                   'retrieval_seconds': 0, 'chat_seconds': 0, 'judge_seconds': 0}
            started = time.monotonic()
            stage, phase_started = 'input', started
            print(f'[{i}/{len(cases)}] {case["sheet"]}/{case["test_case_id"]}', flush=True)
            try:
                if case['input_error']:
                    raise ValueError(case['input_error'])
                stage, phase_started = 'retrieval', time.monotonic()
                payload = call_json(session, args.base_url.rstrip('/') + '/api/retrieval/search',
                    dict(query=case['question'], top_k=args.top_k), args.api_timeout)
                row['contexts'] = parse_contexts(payload)
                row['retrieval_seconds'] = round(time.monotonic() - phase_started, 3)
                stage, phase_started = 'chat', time.monotonic()
                payload = call_json(session, args.base_url.rstrip('/') + '/api/rag/chat',
                    dict(query=case['question'], top_k=args.top_k, top_n=args.top_n,
                         threshold=args.threshold, provider='gemini'), args.api_timeout)
                row['answer'] = parse_answer(payload)
                row['chat_seconds'] = round(time.monotonic() - phase_started, 3)
                stage, phase_started = 'judge', time.monotonic()
                # Giữ context rỗng như API trả; Ragas có thể trả NaN cho metric.
                score_case(row, evaluator)
                row['judge_seconds'] = round(time.monotonic() - phase_started, 3)
            except Exception as exc:
                row['status'] = stage.upper() + '_ERROR'
                row['error'] = safe_error(exc)
                if stage != 'input':
                    row[stage + '_seconds'] = round(time.monotonic() - phase_started, 3)
            except KeyboardInterrupt:
                row['status'], row['error'] = 'INTERRUPTED', 'Người dùng dừng khi đang xử lý ' + stage
                row['total_seconds'] = round(time.monotonic() - started, 3)
                rows.append(row)
                save_outputs(rows, args, metadata)
                raise
            row['total_seconds'] = round(time.monotonic() - started, 3)
            rows.append(row)
            save_outputs(rows, args, metadata)
            print(f'  {row["status"]} · {row["total_seconds"]}s · {row["error"]}', flush=True)
    print(f'Đã ghi {args.output / "results_raw.csv"} và {args.output / "dashboard.html"}')
    return 0 if all(r['status'] == 'OK' for r in rows) else 2


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--input', type=Path, default=Path('RAG_TestCase.xlsx'))
    parser.add_argument('--output', type=Path, default=Path('ragas_results'))
    parser.add_argument('--base-url', default='http://127.0.0.1:8001')
    parser.add_argument('--env-file', type=Path)
    parser.add_argument('--limit', type=int)
    parser.add_argument('--top-k', type=int, default=15)
    parser.add_argument('--top-n', type=int, default=5)
    parser.add_argument('--threshold', type=float, default=0.35)
    parser.add_argument('--api-timeout', type=float, default=360)
    parser.add_argument('--judge-timeout', type=float, default=240)
    parser.add_argument('--judge-model', default='gemini-2.5-flash')
    parser.add_argument('--embedding-model', default='models/gemini-embedding-001')
    parser.add_argument('--overwrite', action='store_true')
    args = parser.parse_args()
    if args.limit is not None and args.limit < 1:
        parser.error('--limit phải >= 1')
    if not 1 <= args.top_n <= args.top_k <= 20 or not 0 <= args.threshold <= 1:
        parser.error('Cần 1 <= top-n <= top-k <= 20 và 0 <= threshold <= 1')
    if min(args.api_timeout, args.judge_timeout) <= 0:
        parser.error('Timeout phải > 0')
    try:
        return run(args)
    except KeyboardInterrupt:
        print('Đã dừng; kết quả đã xử lý được lưu.', file=sys.stderr)
        return 130
    except Exception as exc:
        import traceback
        traceback.print_exc()
        print('Không thể khởi động: ' + safe_error(exc), file=sys.stderr)
        return 1


if __name__ == '__main__':
    sys.exit(main())
