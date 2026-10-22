#!/usr/bin/env python3
"""
End-to-end test run of TC_LuatLD.xlsx against the running API, with an Excel report.

1. Calls POST /api/rag/chat for every test case (server .env config, Gemini).
   Raw responses go to <out-dir>/raw.jsonl after each case; re-running the same
   command resumes and skips cases already done (quota / network interruptions).
2. Auto-scores each answer (keywords, expected-answer recall, cited source, page,
   temporal caveat) and suggests PASS / PARTIAL / FAIL / BLOCKED.
3. Writes <out-dir>/TC_LuatLD_results.xlsx: Results (tester fills Final Status),
   Summary (formulas follow Final Status), Retrieval / Regression (optional
   run_tc_retrieval.py JSON), Defects (candidates), Environment.

    # API must be running: uvicorn src.app:app --port 8001
    python evaluation/run_tc_excel.py --input ~/Downloads/TC_LuatLD.xlsx \
        --out-dir evaluation/results/2026-10-20 \
        --retrieval evaluation/results/retrieval_new.json \
        --retrieval-baseline evaluation/results/retrieval_old.json
    # Rebuild the Excel only (no API calls):
    python evaluation/run_tc_excel.py ... --report-only
"""
import argparse
import json
import os
import subprocess
import sys
import time
import urllib.error
import urllib.request
from collections import Counter, OrderedDict, defaultdict
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))

from tc_common import (  # noqa: E402
    GROUP_DOCS, answer_recall, keyword_coverage, load_cases, norm, same_doc, sha256,
)

NO_INFO = "không tìm thấy thông tin để trả lời"
CAVEAT_WORDS = ["hiệu lực", "thời điểm", "cập nhật", "sửa đổi", "hiện hành", "văn bản mới", "thay thế"]
QUOTA_MARKERS = ["429", "RESOURCE_EXHAUSTED", "quota", "503", "UNAVAILABLE", "overloaded"]
CONFIG_KEYS = ["EMBEDDING_MODEL", "RERANKER_MODEL", "RETRIEVAL_TOP_K", "RERANK_TOP_N", "RERANK_THRESHOLD",
               "MAX_CONTEXT_TOKENS", "LLM_PROVIDER", "LLM_MODEL", "GEMINI_MODEL", "OPENAI_MODEL", "OLLAMA_MODEL",
               "PREFERRED_SOURCES", "TORCH_DEVICE"]
GROUP_NAMES = {g: Path(d).stem for g, d in GROUP_DOCS.items()}


# ----------------------------------------------------------------------------
# 1. Run
# ----------------------------------------------------------------------------

def post_json(url, body, timeout):
    req = urllib.request.Request(url, json.dumps(body).encode("utf-8"), {"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return json.loads(resp.read().decode("utf-8"))


def call_chat(api, query, provider, timeout, retries):
    """Returns (response | None, error | None, attempts). Retries quota/overload errors with backoff."""
    body = {"query": query} | ({"provider": provider} if provider else {})
    delay = 30
    for attempt in range(1, retries + 2):
        try:
            return post_json(f"{api}/api/rag/chat", body, timeout), None, attempt
        except urllib.error.HTTPError as exc:
            detail = exc.read().decode("utf-8", "replace")
            error = f"HTTP {exc.code}: {detail[:500]}"
        except Exception as exc:  # network / timeout
            error = f"{type(exc).__name__}: {exc}"
        if attempt > retries or not any(m.lower() in error.lower() for m in QUOTA_MARKERS + ["timed out", "Connection"]):
            return None, error, attempt
        print(f"    retry {attempt}/{retries} in {delay}s: {error[:120]}", flush=True)
        time.sleep(delay)
        delay = min(delay * 2, 300)
    return None, error, attempt


def run_cases(cases, args, raw_path):
    done = {}
    if raw_path.exists():
        for line in raw_path.read_text(encoding="utf-8").splitlines():
            if line.strip():
                rec = json.loads(line)
                if rec.get("response") is not None or not args.retry_errors:
                    done[rec["id"]] = rec
    todo = [c for c in cases if c["id"] not in done]
    print(f"{len(done)} cases already recorded, {len(todo)} to run", flush=True)

    quota_streak = 0
    with raw_path.open("a", encoding="utf-8") as out:
        for i, case in enumerate(todo, 1):
            started = time.time()
            response, error, attempts = call_chat(args.api, case["query"], args.provider, args.timeout, args.retries)
            rec = {"id": case["id"], "query": case["query"], "response": response, "error": error,
                   "attempts": attempts, "client_seconds": round(time.time() - started, 2),
                   "run_at": datetime.now().isoformat(timespec="seconds")}
            out.write(json.dumps(rec, ensure_ascii=False) + "\n")
            out.flush()
            done[case["id"]] = rec
            status = "OK" if response else f"ERROR {error[:100]}"
            print(f"[{i}/{len(todo)}] {case['id']} {rec['client_seconds']}s {status}", flush=True)

            quota_streak = quota_streak + 1 if error and any(m.lower() in error.lower() for m in QUOTA_MARKERS) else 0
            if quota_streak >= 3:
                print("Stopping: 3 quota errors in a row. Re-run the same command later to resume.", flush=True)
                break
            if args.sleep:
                time.sleep(args.sleep)
    return done


# ----------------------------------------------------------------------------
# 2. Score
# ----------------------------------------------------------------------------

def score(case, rec):
    resp = (rec or {}).get("response")
    if not resp:
        return {"status": "BLOCKED", "reasons": [f"Lỗi API: {(rec or {}).get('error') or 'chưa chạy'}"], "answer": "",
                "cited": [], "context": [], "latency": None, "model": ""}

    answer = resp.get("answer") or ""
    cited = resp.get("citations") or []
    context = (resp.get("debug") or {}).get("context_sources") or cited
    refused = NO_INFO in norm(answer)
    kw = keyword_coverage(case["keywords"], answer)
    recall = answer_recall(case["expected"], answer)
    cited_doc = any(same_doc(c.get("file"), case["expected_doc"]) for c in cited)
    context_doc = any(same_doc(c.get("file"), case["expected_doc"]) for c in context)
    page_ok = None
    if case["page"] is not None:
        page_ok = any(same_doc(c.get("file"), case["expected_doc"]) and abs(int(c.get("page") or 0) - case["page"]) <= 1 for c in cited)
    needs_caveat = "temporal" in case["behavior"].lower()
    caveat_ok = any(w in norm(answer) for w in CAVEAT_WORDS) if needs_caveat else None

    reasons = []
    answerable = case["answerable"].lower() != "no"
    leaked = [k for k in case.get("forbidden") or [] if norm(k) in norm(answer)]
    if leaked:
        status = "FAIL"
        reasons.append("Có nội dung bị cấm: " + "; ".join(leaked))
    elif refused:
        status = "FAIL" if answerable else "PASS"
        reasons.append("Từ chối trả lời dù câu hỏi có đáp án" if answerable else "Từ chối đúng")
    elif not answerable and str(case.get("category") or "").lower() == "security":
        # A safe answer may decline in its own words; the tester decides.
        status = "PARTIAL"
        reasons.append("Không dùng câu từ chối chuẩn - tester kiểm tra có làm theo yêu cầu tấn công không")
    elif not answerable:
        status = "FAIL"
        reasons.append("Trả lời câu hỏi không có đáp án (nguy cơ bịa)")
    else:
        if kw is not None and kw < 0.7:
            reasons.append(f"Thiếu từ khóa ({kw:.0%})")
        if recall is not None and recall < 0.5:
            reasons.append(f"Ít khớp Expected Answer ({recall:.0%})")
        if not cited_doc:
            reasons.append("Không trích tài liệu mong đợi" + (" (có trong ngữ cảnh)" if context_doc else " (không có trong ngữ cảnh)"))
        if page_ok is False:
            reasons.append(f"Không trích đúng trang {case['page']} (±1)")
        if caveat_ok is False:
            reasons.append("Thiếu lưu ý về thời điểm hiệu lực")
        weak = (kw is not None and kw < 0.3) and (recall is not None and recall < 0.3)
        status = "FAIL" if weak else ("PASS" if not reasons or reasons == [f"Không trích đúng trang {case['page']} (±1)"] else "PARTIAL")

    llm_log = next((l for l in resp.get("logs") or [] if l.get("stage") == "llm_generation"), {})
    model = llm_log.get("model") or llm_log.get("provider") or ("(no LLM call: early exit)" if refused else "")

    return {
        "status": status, "reasons": reasons, "answer": answer, "cited": cited, "context": context, "model": model,
        "refused": refused, "kw": kw, "recall": recall, "cited_doc": cited_doc, "context_doc": context_doc,
        "page_ok": page_ok, "caveat_ok": caveat_ok,
        "latency": round((resp.get("total_duration_ms") or 0) / 1000, 1),
    }


def fmt_sources(items):
    return "\n".join(f"[{c.get('source_id', '')}] {Path(c.get('file') or '').stem} | tr.{c.get('page')} | "
                     f"{str(c.get('section') or '')[:60]} | {c.get('rerank_score')}" for c in items)


# ----------------------------------------------------------------------------
# 3. Excel report
# ----------------------------------------------------------------------------

def git_info():
    try:
        head = subprocess.run(["git", "rev-parse", "--short", "HEAD"], cwd=ROOT, capture_output=True, text=True).stdout.strip()
        dirty = subprocess.run(["git", "status", "--porcelain", "src", "frontend/src", "evaluation"], cwd=ROOT,
                               capture_output=True, text=True).stdout.strip().splitlines()
        return head, len(dirty)
    except Exception:
        return "unknown", 0


MANUAL_COLUMNS = ("Final Status", "Defect ID", "Tester Note")


def load_manual(out_xlsx):
    """Reviewer columns from a previous report, so --report-only never wipes a review."""
    if not Path(out_xlsx).exists():
        return {}
    from openpyxl import load_workbook
    ws = load_workbook(out_xlsx, read_only=True)["Results"]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    if not all(c in header for c in MANUAL_COLUMNS):
        return {}
    idx = [header.index(c) for c in MANUAL_COLUMNS]
    out = {}
    for r in rows:
        vals = tuple(r[i] for i in idx)
        if any(v not in (None, "") for v in vals):
            out[str(r[header.index("Test Case ID")])] = vals
    return out


def build_report(header, cases, raw, args, out_xlsx):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(bold=True, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")
    fills = {"PASS": "C6EFCE", "PARTIAL": "FFEB9C", "FAIL": "FFC7CE", "BLOCKED": "D9D9D9"}

    def style_header(ws, row=1):
        for cell in ws[row]:
            cell.fill, cell.font, cell.alignment = head_fill, head_font, Alignment(wrap_text=True, vertical="center")

    scored = OrderedDict((c["id"], score(c, raw.get(c["id"]))) for c in cases)
    manual = load_manual(out_xlsx)
    wb = Workbook()

    # --- Results ---
    ws = wb.active
    ws.title = "Results"
    extra = ["Document", "Actual Answer", "Cited Sources", "Context Sources", "LLM Model", "Latency (s)", "Keyword Coverage",
             "Expected-Answer Recall", "Cited Expected Doc", "Page Match", "Suggested Status", "Auto Reasons",
             "Final Status", "Effective Status", "Defect ID", "Tester Note"]
    cols = header + extra
    ws.append(cols)
    style_header(ws)
    col = {name: i + 1 for i, name in enumerate(cols)}
    L = {name: get_column_letter(i) for name, i in col.items()}
    for r, case in enumerate(cases, start=2):
        s = scored[case["id"]]
        final, defect, note = manual.get(case["id"], (None, None, None))
        values = [case["row"].get(h) for h in header] + [
            GROUP_NAMES.get(case["group"], ""), s["answer"], fmt_sources(s["cited"]), fmt_sources(s["context"]),
            s["model"], s["latency"], s.get("kw"), s.get("recall"),
            None if s["status"] == "BLOCKED" else ("Yes" if s.get("cited_doc") else "No"),
            {True: "Yes", False: "No", None: "n/a"}[s.get("page_ok")],
            s["status"], "; ".join(s["reasons"]), final, None, defect, note,
        ]
        ws.append(values)
        ws.cell(r, col["Effective Status"]).value = f'=IF({L["Final Status"]}{r}="",{L["Suggested Status"]}{r},{L["Final Status"]}{r})'
        for name in ("Keyword Coverage", "Expected-Answer Recall"):
            ws.cell(r, col[name]).number_format = "0%"
    last = len(cases) + 1
    for row in ws.iter_rows(min_row=2, max_row=last):
        for cell in row:
            cell.alignment = wrap
    widths = {"Query": 45, "Expected Answer": 45, "Evidence": 35, "Actual Answer": 60, "Cited Sources": 45,
              "Context Sources": 45, "Auto Reasons": 35, "Tester Note": 30, "Required Keywords / Key Concepts": 25}
    for name, i in col.items():
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 14)
    ws.freeze_panes = ws.cell(2, col["Query"] + 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
    dv = DataValidation(type="list", formula1='"PASS,PARTIAL,FAIL,BLOCKED"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{L['Final Status']}2:{L['Final Status']}{last}")
    for name in ("Suggested Status", "Final Status", "Effective Status"):
        rng = f"{L[name]}2:{L[name]}{last}"
        for status, color in fills.items():
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{status}"'], fill=PatternFill("solid", fgColor=color)))

    # --- Summary (formulas over Effective Status so tester edits update it) ---
    sm = wb.create_sheet("Summary", 0)
    eff = f"Results!${L['Effective Status']}$2:${L['Effective Status']}${last}"
    statuses = ["PASS", "PARTIAL", "FAIL", "BLOCKED"]
    sm.append(["TC_LuatLD — Test Summary"]); sm["A1"].font = Font(bold=True, size=14)
    sm.append(["Effective Status = Final Status (tester) nếu đã điền, nếu không = Suggested Status (tự động)."])
    sm.append([])
    sm.append(["Status", "Count", "Rate"]); style_header(sm, sm.max_row)
    first = sm.max_row + 1
    for st in statuses:
        r = sm.max_row + 1
        sm.append([st, f'=COUNTIF({eff},"{st}")', f"=B{r}/{len(cases)}"])
        sm[f"C{r}"].number_format = "0.0%"
    sm.append(["Total", f"=SUM(B{first}:B{sm.max_row})"])
    sm.append(["Reviewed by tester", f'=COUNTA(Results!{L["Final Status"]}2:{L["Final Status"]}{last})'])
    lat = [s["latency"] for s in scored.values() if s["latency"]]
    sm.append(["Median latency (s)", sorted(lat)[len(lat) // 2] if lat else None])

    def breakdown(title, field_name):
        sm.append([]); sm.append([title]); sm.cell(sm.max_row, 1).font = bold
        sm.append([field_name, "Total"] + statuses + ["Pass rate"]); style_header(sm, sm.max_row)
        values = list(OrderedDict.fromkeys(str(c["row"].get(field_name) if field_name != "Document" else GROUP_NAMES.get(c["group"], "")) for c in cases))
        key_col = f"Results!${L[field_name]}$2:${L[field_name]}${last}"
        for v in values:
            r = sm.max_row + 1
            row = [v, f'=COUNTIF({key_col},A{r})']
            row += [f'=COUNTIFS({key_col},A{r},{eff},"{st}")' for st in statuses]
            row += [f"=IF(B{r}=0,0,C{r}/B{r})"]
            sm.append(row)
            sm[f"H{r}"].number_format = "0%"

    breakdown("Theo tài liệu", "Document")
    breakdown("Theo độ khó", "Retrieval Difficulty")
    breakdown("Theo loại câu hỏi", "Category")
    sm.column_dimensions["A"].width = 48
    for c in "BCDEFGH":
        sm.column_dimensions[c].width = 11

    # --- Retrieval / Regression ---
    def load_retrieval(path):
        return json.loads(Path(path).read_text(encoding="utf-8")) if path else None

    def retrieval_agg(rows):
        n = len(rows) or 1
        pages = [r["page_hit_at5"] for r in rows if r["page_hit_at5"] is not None]
        kws = [r["keyword_coverage"] for r in rows if r["keyword_coverage"] is not None]
        return [sum(r["doc_hit_at1"] for r in rows) / n, sum(r["doc_hit_at5"] for r in rows) / n,
                sum(pages) / len(pages) if pages else None, sum(kws) / len(kws) if kws else None,
                sum(not r["passed_threshold"] for r in rows) / n]

    new_ret, old_ret = load_retrieval(args.retrieval), load_retrieval(args.retrieval_baseline)
    if new_ret:
        rt = wb.create_sheet("Retrieval")
        rt.append(["ID", "Document", "Category", "Difficulty", "Doc@1", "Doc@5", "Page@5 (±1)", "Keyword coverage (top 5)",
                   "Best rerank", "Refused (< threshold)", "Top sources"]); style_header(rt)
        for r in new_ret["results"]:
            rt.append([r["id"], GROUP_NAMES.get(r["group"]), r["category"], r["difficulty"], int(r["doc_hit_at1"]), int(r["doc_hit_at5"]),
                       r["page_hit_at5"], r["keyword_coverage"], r.get("best_rerank"), int(not r["passed_threshold"]),
                       "\n".join(f"{Path(t['source']).stem} tr.{t['pages'][0]}-{t['pages'][1]} ({t['rerank']})" for t in r["top"])])
        rt.column_dimensions["K"].width = 60
        rt.column_dimensions["B"].width = 36

        rg = wb.create_sheet("Regression")
        metrics = ["Doc@1", "Doc@5", "Page@5", "Keywords@5", "Refused"]
        rg.append(["Retrieval metrics" + (f": {old_ret['label']} vs {new_ret['label']}" if old_ret else "")]); rg["A1"].font = Font(bold=True, size=13)
        rg.append(["Scope"] + [f"{m} {lab}" for m in metrics for lab in ([old_ret["label"], new_ret["label"]] if old_ret else [new_ret["label"]])])
        style_header(rg, 2)
        scopes = [("ALL", lambda r: True)] + [(GROUP_NAMES[g], (lambda g: lambda r: r["group"] == g)(g)) for g in GROUP_DOCS] + \
                 [(d, (lambda d: lambda r: r["difficulty"] == d)(d)) for d in ("Easy", "Medium", "Hard")]
        for name, pred in scopes:
            a_new = retrieval_agg([r for r in new_ret["results"] if pred(r)])
            a_old = retrieval_agg([r for r in old_ret["results"] if pred(r)]) if old_ret else None
            row = [name]
            for i in range(len(metrics)):
                row += [a_old[i], a_new[i]] if old_ret else [a_new[i]]
            rg.append(row)
            for cell in rg[rg.max_row][1:]:
                cell.number_format = "0%"
        rg.column_dimensions["A"].width = 44
        if old_ret:
            o = {r["id"]: r for r in old_ret["results"]}
            fixed = [r["id"] for r in new_ret["results"] if r["doc_hit_at5"] and not o[r["id"]]["doc_hit_at5"]]
            broken = [r["id"] for r in new_ret["results"] if o[r["id"]]["doc_hit_at5"] and not r["doc_hit_at5"]]
            rg.append([]); rg.append([f"Doc@5 fixed ({len(fixed)})", ", ".join(fixed)])
            rg.append([f"Doc@5 regressed ({len(broken)})", ", ".join(broken)])

    # --- Defects (candidates from FAIL / BLOCKED; tester confirms) ---
    df = wb.create_sheet("Defects")
    df.append(["Defect ID", "Test Case", "Severity", "Title", "Steps to reproduce", "Expected", "Actual", "Auto reasons", "Status"])
    style_header(df)
    n = 0
    for case in cases:
        s = scored[case["id"]]
        if s["status"] not in ("FAIL", "BLOCKED"):
            continue
        n += 1
        severity = "High" if (s.get("refused") and case["difficulty"] == "Easy") or s["status"] == "BLOCKED" else "Medium"
        title = "Từ chối trả lời câu hỏi có đáp án" if s.get("refused") else ("Lỗi API" if s["status"] == "BLOCKED" else "Câu trả lời sai/thiếu nội dung")
        df.append([f"D-{n:03d}", case["id"], severity, title,
                   f'1. Mở RAG Lab / POST /api/rag/chat\n2. Query: "{case["query"]}"',
                   case["expected"], s["answer"][:1500], "; ".join(s["reasons"]), "Open (candidate)"])
    for c, w in zip("ABCDEFGHI", (10, 12, 10, 30, 45, 45, 55, 35, 16)):
        df.column_dimensions[c].width = w
    for row in df.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap

    # --- Environment ---
    env = wb.create_sheet("Environment")
    head, dirty = git_info()
    try:
        status = urllib.request.urlopen(f"{args.api}/api/pipeline/status", timeout=30).read().decode("utf-8")
        status = json.loads(status)
    except Exception:
        status = {}
    runs = [r.get("run_at") for r in raw.values() if r.get("run_at")]
    rows = [
        ("Report generated", datetime.now().isoformat(timespec="seconds")),
        ("Test execution", f"{min(runs)} → {max(runs)}" if runs else ""),
        ("Test file", f"{args.input.name} (sha256 {sha256(args.input)[:16]})"),
        ("Test cases", len(cases)),
        ("API", args.api),
        ("LLM provider override", args.provider or "(server default)"),
        ("LLM models used (from responses)", "; ".join(f"{m or '(none)'}: {n}" for m, n in
                                                        Counter(s["model"] for s in scored.values()).most_common())),
        ("Git commit", f"{head} + {dirty} uncommitted file(s) in src/, frontend/src, evaluation/" if dirty else head),
        ("Indexed documents", status.get("pdf_documents")),
        ("Pages / elements / chunks / vectors", f"{status.get('pages')} / {status.get('elements')} / {status.get('chunks')} / {status.get('vectors')}"),
        ("Embedding model / device", f"{status.get('embedding_model')} / {status.get('device')}"),
    ]
    sys.path.insert(0, str(ROOT))
    from src.env import parse_env_file  # local .env, secrets excluded
    dotenv = parse_env_file(ROOT / ".env")
    rows += [(f"config {k}", os.environ.get(k) or dotenv.get(k, "(default)")) for k in CONFIG_KEYS]
    rows += [("Group → document", "; ".join(f"{g}: {d}" for g, d in GROUP_DOCS.items()))]
    env.append(["Item", "Value"]); style_header(env)
    for row in rows:
        env.append(list(row))
    env.column_dimensions["A"].width = 36
    env.column_dimensions["B"].width = 110

    add_eval_sheets(wb, args.out_dir, style_header, wrap)
    wb.save(out_xlsx)
    counts = OrderedDict((st, sum(s["status"] == st for s in scored.values())) for st in statuses)
    return counts


def add_eval_sheets(wb, out_dir, style_header, wrap):
    """Sheets from the other evaluation outputs found in out_dir (each one optional)."""
    from openpyxl.styles import Font

    def pct(ws, cols):
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if cell.column_letter in cols and isinstance(cell.value, float):
                    cell.number_format = "0.0%" if cols.get(cell.column_letter) == "%" else "0.00"

    def mean(xs):
        xs = [x for x in xs if x is not None]
        return sum(xs) / len(xs) if xs else None

    recall_path = out_dir / "recall_at5.json"
    recall = json.loads(recall_path.read_text(encoding="utf-8")) if recall_path.exists() else None
    if recall:
        s = recall["summary"]
        ws = wb.create_sheet("Recall@5")
        ws.append(["Recall@5 trên gold chunk (gán tự động, tester xác nhận)"]); ws["A1"].font = Font(bold=True, size=13)
        ws.append(["Scope", "Cases", "With gold", "Recall@5 (context to LLM)", "Recall@5 (ranking)", "Hit@5", "MRR@5", "Gold in hybrid candidates"])
        style_header(ws, 2)
        scopes = [("ALL", s["all"])] + [(GROUP_NAMES.get(k, k), v) for k, v in s["by_group"].items()] + list(s["by_difficulty"].items())
        for name, a in scopes:
            ws.append([name, a["n"], a["with_gold"], a["recall_at5_context"], a["recall_at5_ranked"], a["hit_at5_ranked"],
                       a["mrr_at5"], a["gold_in_candidates"]])
        ws.append([])
        gold_path = out_dir / "gold_review.json"
        gold_review = json.loads(gold_path.read_text(encoding="utf-8")) if gold_path.exists() else {}
        ws.append(["ID", "Group", "Category", "Difficulty", "Status", "Stage", "Gold chunks", "Gold score",
                   "Recall@5 ctx", "Recall@5 rank", "RR@5", "Gold in candidates", "Gold review"])
        style_header(ws, ws.max_row)
        for r in recall["results"]:
            review = gold_review.get("cases", {}).get(r["id"]) or (gold_review.get("default") if r["gold"] else None)
            ws.append([r["id"], r["group"], r["category"], r["difficulty"], r["status"], r["stage"], "\n".join(r["gold"]),
                       r["gold_score"], r["recall_at5_context"], r["recall_at5_ranked"], r["rr_at5"], r["gold_in_candidates"], review])
        pct(ws, {c: "%" for c in "DEFHIJ"})
        ws.column_dimensions["A"].width = 40
        ws.column_dimensions["G"].width = 50

    ragas_path = out_dir / "ragas.jsonl"
    ragas = {}
    if ragas_path.exists():
        for line in ragas_path.read_text(encoding="utf-8").splitlines():
            if line.strip():
                rec = json.loads(line)
                ragas[rec["id"]] = rec  # last record wins (resumed runs)
    metrics = ("faithfulness", "answer_relevancy", "context_precision", "context_recall")
    if ragas:
        scored = [r for r in ragas.values() if r["status"] in ("OK", "PARTIAL_METRICS", "NO_CONTEXT")]
        ws = wb.create_sheet("RAGAS")
        ws.append([f"RAGAS ({len(scored)} câu đã chấm; judge {next(iter(ragas.values())).get('judge_model')})"])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append(["Faithfulness của câu từ chối = 0 theo RAGAS; dòng 'answered, not refused' loại các câu đó."])
        ws.append(["Scope", "n"] + list(metrics)); style_header(ws, 3)
        stage = {r["id"]: r["stage"] for r in (recall or {}).get("results", [])}
        refused = {r["id"] for r in (recall or {}).get("results", []) if r["status"] == "FAIL"}
        scopes = [("ALL scored", scored), ("answered, not refused", [r for r in scored if r["status"] != "NO_CONTEXT" and r["id"] not in refused])]
        scopes += [(GROUP_NAMES.get(g, g), [r for r in scored if r["group"] == g]) for g in GROUP_DOCS]
        scopes += [(d, [r for r in scored if r["difficulty"] == d]) for d in ("Easy", "Medium", "Hard")]
        scopes += [(f"stage {st}", [r for r in scored if stage.get(r["id"]) == st]) for st in sorted(set(stage.values()))]
        for name, rs in scopes:
            ws.append([name, len(rs)] + [mean(r.get(m) for r in rs) for m in metrics])
        ws.append([])
        ws.append(["ID", "Group", "Difficulty", "Status", "Stage"] + list(metrics) + ["Errors"]); style_header(ws, ws.max_row)
        for r in sorted(ragas.values(), key=lambda r: r["id"]):
            ws.append([r["id"], r["group"], r["difficulty"], r["status"], stage.get(r["id"])] +
                      [r.get(m) for m in metrics] + ["; ".join(r.get("errors") or [])[:300]])
        pct(ws, {c: "n" for c in "CDEFGHI"})
        ws.column_dimensions["A"].width = 32

    if recall:
        ws = wb.create_sheet("Failure Analysis")
        ws.append(["Phân tích sai lệch theo tầng (cây quyết định mục 7 quy trình QA)"]); ws["A1"].font = Font(bold=True, size=13)
        ws.append(["Stage", "Meaning", "Cases", "Share", "IDs"]); style_header(ws, 2)
        labels = recall["summary"]["stage_labels"]
        by = defaultdict(list)
        for r in recall["results"]:
            by[r["stage"]].append(r["id"])
        for st, ids in sorted(by.items(), key=lambda x: -len(x[1])):
            ws.append([st, labels.get(st, ""), len(ids), len(ids) / len(recall["results"]), ", ".join(ids)])
        pct(ws, {"D": "%"})
        for c, w in zip("ABCDE", (24, 60, 8, 8, 90)):
            ws.column_dimensions[c].width = w
        for row in ws.iter_rows(min_row=3):
            for cell in row:
                cell.alignment = wrap

    api_path = out_dir / "api_tests.json"
    if api_path.exists():
        ws = wb.create_sheet("API Tests")
        ws.append(["ID", "Test", "Request", "Expected", "HTTP status", "Seconds", "Result", "LLM provider", "Response (excerpt)"])
        style_header(ws)
        for r in json.loads(api_path.read_text(encoding="utf-8"))["results"]:
            ws.append([r["id"], r["title"], r["request"], r["expected"], r["status"], r["seconds"], r["result"], r["llm_provider"], r["body"][:300]])
        for c, w in zip("ABCDEFGHI", (9, 34, 45, 40, 8, 8, 8, 14, 60)):
            ws.column_dimensions[c].width = w

    static_path = out_dir / "static_checks.json"
    if static_path.exists():
        st = json.loads(static_path.read_text(encoding="utf-8"))
        ws = wb.create_sheet("Static Checks")
        ws.append(["Group 1 - Ingestion"]); ws["A1"].font = Font(bold=True)
        ws.append(["Document", "PDF pages", "Pages with text missing", "Chunks", "Fragment rate /1000 words", "Odd chars", "Junk chunks", "Repeated lines (watermark)"])
        style_header(ws, 2)
        for doc, d in st["ingestion"].items():
            ws.append([doc, d["pdf_pages"], d["pages_missing"], d["chunks"], d["fragment_rate"], d["odd_chars"], d["junk_chunks"],
                       "\n".join(f"x{r['chunks']}: {r['line']}" for r in d["repeated_lines"])])
        ch = st["chunking"]
        ws.append([]); ws.append(["Group 2 - Chunking", f"{ch['chunks']} chunks, {ch['unique_ids']} unique ids, tokens p50 {ch['tokens_p50']}, max {ch['tokens_max']}"])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for k, v in ch["issues"].items():
            ws.append([k, v["count"], ", ".join(v["examples"])])
        c = st["citations"]
        ws.append([]); ws.append(["Group 5 - Citations", f"answered {c['answered']}, with >= 1 citation {c['with_citation']}, cites expected doc {c['cites_expected_doc']}"])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for k in ("no_citation", "combined_tag_cases", "dropped_tag_cases", "invented_tag_cases"):
            ws.append([k, len(c[k]), ", ".join(c[k])])
        for col, w in zip("ABCDEFGH", (46, 14, 14, 10, 14, 10, 10, 50)):
            ws.column_dimensions[col].width = w


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", type=Path, required=True, help="TC_LuatLD.xlsx")
    ap.add_argument("--out-dir", type=Path, required=True)
    ap.add_argument("--api", default="http://127.0.0.1:8001")
    ap.add_argument("--provider", default=None, help="gemini | openai | ollama | mock (default: server LLM_PROVIDER)")
    ap.add_argument("--limit", type=int, default=None, help="only the first N cases")
    ap.add_argument("--ids", default=None, help="comma-separated test case IDs")
    ap.add_argument("--timeout", type=int, default=600)
    ap.add_argument("--retries", type=int, default=4)
    ap.add_argument("--sleep", type=float, default=0, help="pause between cases (rate limits)")
    ap.add_argument("--retry-errors", action="store_true", help="re-run cases recorded with an error")
    ap.add_argument("--retrieval", type=Path, default=None, help="run_tc_retrieval.py JSON for this build")
    ap.add_argument("--retrieval-baseline", type=Path, default=None, help="run_tc_retrieval.py JSON of a previous build")
    ap.add_argument("--report-only", action="store_true")
    args = ap.parse_args()
    args.input = args.input.expanduser()

    header, cases = load_cases(args.input)
    if args.ids:
        wanted = {i.strip() for i in args.ids.split(",")}
        cases = [c for c in cases if c["id"] in wanted]
    if args.limit:
        cases = cases[: args.limit]
    args.out_dir.mkdir(parents=True, exist_ok=True)
    raw_path = args.out_dir / "raw.jsonl"

    if args.report_only:
        raw = {}
        if raw_path.exists():
            for line in raw_path.read_text(encoding="utf-8").splitlines():
                if line.strip():
                    rec = json.loads(line)
                    raw[rec["id"]] = rec
    else:
        raw = run_cases(cases, args, raw_path)

    out_xlsx = args.out_dir / f"{args.input.stem}_results.xlsx"
    counts = build_report(header, cases, raw, args, out_xlsx)
    print(f"Suggested: {dict(counts)}")
    print(f"Report: {out_xlsx}")


if __name__ == "__main__":
    main()
